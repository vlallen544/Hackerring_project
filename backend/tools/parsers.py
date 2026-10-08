# Parsers: turn source files (PDF, PPTX, MD, TXT) into page-level chunks that keep source id + page number
import hashlib
import json
import re
import unicodedata
from collections import Counter
from dataclasses import asdict, dataclass
from pathlib import Path

SUPPORTED_SUFFIXES = {".pdf", ".pptx", ".md", ".txt"}


@dataclass
class Chunk:
    source_id: str
    page: int
    text: str

    def to_dict(self):
        return asdict(self)


def file_sha256(path):
    """Fingerprint of a source file, used to tell whether it changed since the last course build."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(1 << 20), b""):
            h.update(block)
    return h.hexdigest()


def parse_page_range(spec, total):
    """'1-5, 8' -> [1, 2, 3, 4, 5, 8] (1-based, within 1..total). Empty spec means all pages."""
    if not spec or not str(spec).strip():
        return list(range(1, total + 1))
    pages = set()
    for part in str(spec).replace(" ", "").split(","):
        m = re.fullmatch(r"(\d+)(?:-(\d+))?", part)
        if not m:
            raise ValueError(f"Invalid page range '{spec}'. Use a form like 12-40 or 1-5,8")
        start, end = int(m.group(1)), int(m.group(2) or m.group(1))
        if start < 1 or end < start or end > total:
            raise ValueError(f"Page range '{part}' is outside 1-{total}")
        pages.update(range(start, end + 1))
    return sorted(pages)


# --------------------------------------------------------------------------- #
# PDF text cleanup: pdfplumber returns layout text, so fix what breaks exact-quote verification
# --------------------------------------------------------------------------- #
_PAGE_NUMBER = re.compile(r"^\s*(page\s*)?\d+(\s*(of|/)\s*\d+)?\s*$", re.IGNORECASE)


def _clean_pdf_pages(raw_pages):
    """raw_pages: list of page texts. Removes running headers/footers and page numbers, joins hyphenated words."""
    lines_per_page = [[ln.strip() for ln in unicodedata.normalize("NFKC", t).splitlines() if ln.strip()]
                      for t in raw_pages]
    # A line that appears at the top or bottom of at least half of the pages is a running header/footer
    repeated = set()
    if len(lines_per_page) >= 2:
        edges = Counter()
        for lines in lines_per_page:
            edges.update({re.sub(r"\d+", "#", ln) for ln in lines[:2] + lines[-2:]})
        repeated = {ln for ln, n in edges.items() if n >= max(2, len(lines_per_page) / 2)}

    cleaned = []
    for lines in lines_per_page:
        kept = [ln for ln in lines if not _PAGE_NUMBER.match(ln) and re.sub(r"\d+", "#", ln) not in repeated]
        text = "\n".join(kept)
        text = re.sub(r"(\w)-\n([a-z])", r"\1\2", text)  # "normal-\nization" -> "normalization"
        cleaned.append(text.strip())
    return cleaned


def _read_pdf(path):
    import pdfplumber

    with pdfplumber.open(path) as pdf:
        raw = [(page.extract_text() or "") for page in pdf.pages]
    return _clean_pdf_pages(raw)


def _read_pages(path):
    """All pages of a file as cleaned text (1-based page i is item i-1)."""
    path = Path(path)
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        return _read_pdf(path)
    if suffix == ".pptx":
        from pptx import Presentation

        return ["\n".join(s.text_frame.text for s in slide.shapes if s.has_text_frame).strip()
                for slide in Presentation(path).slides]
    if suffix in (".md", ".txt"):
        # a line containing only ---page--- marks a page break
        return [p.strip() for p in path.read_text(encoding="utf-8").split("---page---")]
    raise ValueError(f"Unsupported file type: {suffix}")


def parse_file(path, source_id=None, page_range=None):
    """Chunks for the pages in page_range (e.g. '12-40'; all pages if empty). Page numbers stay the file's own."""
    path = Path(path)
    source_id = source_id or path.stem
    pages = _read_pages(path)
    wanted = parse_page_range(page_range, len(pages))
    chunks = [Chunk(source_id, n, pages[n - 1]) for n in wanted if pages[n - 1]]
    if not chunks and path.suffix.lower() == ".pdf":
        raise ValueError(f"{path.name}: no readable text (scanned PDF?)")
    return chunks


def inspect_file(path, page_range=None):
    """Summary used when a file is uploaded: page count, pages used, words, and pages without text."""
    pages = _read_pages(path)
    wanted = parse_page_range(page_range, len(pages))
    return {
        "pages_total": len(pages),
        "pages_used": len(wanted),
        "pages_without_text": [n for n in wanted if not pages[n - 1]],
        "words": sum(len(pages[n - 1].split()) for n in wanted),
    }


def page_texts(path, page_range=None):
    """[{page, words, text}] for previewing exactly what the agents will read."""
    pages = _read_pages(path)
    return [{"page": n, "words": len(pages[n - 1].split()), "text": pages[n - 1]}
            for n in parse_page_range(page_range, len(pages))]


def load_sources(data_dir="sample_data"):
    """Reads sources.json and parses every listed file. Returns (source_metadata, all_chunks)."""
    data_dir = Path(data_dir)
    meta = json.loads((data_dir / "sources.json").read_text(encoding="utf-8"))
    chunks = []
    for src in meta["sources"]:
        chunks += parse_file(data_dir / src["file"], src["id"], src.get("page_range"))
    return meta, chunks


def chunks_to_prompt(chunks):
    """Formats chunks so the model can see (and cite) exactly which source and page each text is from."""
    return "\n\n".join(f"<<SOURCE {c.source_id} | PAGE {c.page}>>\n{c.text}" for c in chunks)
